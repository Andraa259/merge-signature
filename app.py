import streamlit as st
from docx import Document
from docx.shared import Cm, Pt
from datetime import datetime
import io
import time

# Konfigurasi Halaman
st.set_page_config(page_title="Expert Signature Pro", page_icon="✍️")

def get_indo_date():
    months = {1: "Januari", 2: "Februari", 3: "Maret", 4: "April", 5: "Mei", 6: "Juni", 
              7: "Juli", 8: "Agustus", 9: "September", 10: "Oktober", 11: "November", 12: "Desember"}
    now = datetime.now()
    return f"{now.day} {months[now.month]} {now.year}"

# --- Tampilan Header ---
st.title("✍️ Expert Validation Automator")
st.markdown("""
Selamat datang! Gunakan alat ini untuk mengisi **Tanggal**, **Tanda Tangan**, dan **Nama** secara otomatis pada formulir validasi.
""")

# --- Seksi Informasi & Persyaratan ---
with st.expander("ℹ️ Persyaratan & Tips Hasil Terbaik", expanded=True):
    st.markdown("""
    *   **Background TTD:** Sangat disarankan menggunakan format **PNG Transparan** agar hasil terlihat menyatu dengan dokumen.
    *   **Warna Putih:** Jika terpaksa menggunakan background putih, pastikan pencahayaan foto merata.
    *   **Ukuran Gambar:** Sistem akan otomatis menyesuaikan ukuran TTD menjadi standar dokumen (4.5cm).
    *   **Keamanan:** File diproses secara lokal di sesi ini dan tidak disimpan permanen di server.
    """)

st.divider()

# --- Input Section ---
col1, col2 = st.columns(2)
with col1:
    uploaded_file = st.file_uploader("📂 Upload Formulir Word (.docx)", type=["docx"])
with col2:
    img_file = st.file_uploader("🖋️ Upload Scan TTD Online", type=["png", "jpg", "jpeg"])

expert_name = st.text_input("👤 Nama Lengkap Expert Panelis", placeholder="Contoh: Prof. Dr. Andra Raditya, M.Psi.")

st.divider()

# --- Logic Processing ---
if st.button("🚀 Proses & Sinkronisasi Dokumen"):
    if uploaded_file and img_file and expert_name:
        with st.status("Sedang menyusun dokumen...", expanded=True) as status:
            st.write("Membaca struktur file...")
            doc = Document(uploaded_file)
            TTD_WIDTH = Cm(4.5)
            
            st.write("Menyelaraskan tanggal dan posisi TTD...")
            time.sleep(0.5) # Memberi kesan proses elegan
            
            # Kumpulkan semua paragraf (Body + Tabel)
            all_paragraphs = list(doc.paragraphs)
            for table in doc.tables:
                for row in table.rows:
                    for cell in row.cells:
                        all_paragraphs.extend(cell.paragraphs)

            # Proses Manipulasi
            for p in all_paragraphs:
                # LOGIKA A: Baris Surabaya
                if "Surabaya," in p.text:
                    p.text = "" 
                    run = p.add_run(f"Surabaya, {get_indo_date()}")
                    run.add_break() 
                    run.add_picture(img_file, width=TTD_WIDTH)
                    run.add_break() 
                    run.add_text(f"{expert_name}")
                    
                    p.paragraph_format.line_spacing = 1.0
                    p.paragraph_format.space_after = Pt(0)
                    p.paragraph_format.space_before = Pt(0)

                # LOGIKA B: Hapus baris lama secara aman
                elif "(Tt Expert Judgement)" in p.text:
                    try:
                        p_element = p._element
                        parent = p_element.getparent()
                        if parent is not None:
                            parent.remove(p_element)
                    except:
                        p.text = ""
                
                # LOGIKA C: Pembersihan spasi liar
                elif p.text.strip() == "" and "(Tt Expert Judgement)" not in p.text:
                    p.paragraph_format.line_spacing = Pt(1)
                    p.paragraph_format.space_after = Pt(0)
                    p.paragraph_format.space_before = Pt(0)

            # Simpan ke Memory
            target_stream = io.BytesIO()
            doc.save(target_stream)
            
            status.update(label="✅ Dokumen siap diunduh!", state="complete", expanded=False)

        st.balloons()
        st.success("Berhasil! Silakan klik tombol di bawah untuk mengambil file Anda.")
        
        st.download_button(
            label="📥 Download File Hasil Validasi",
            data=target_stream.getvalue(),
            file_name=f"Form Validasi Expert Judgement Forgiveness_{expert_name}.docx",
            mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            use_container_width=True
        )
    else:
        st.error("⚠️ Mohon lengkapi seluruh data (File, TTD, dan Nama) sebelum memproses.")

# --- Footer ---
st.markdown("---")
st.caption("Expert Signature Automator v2.0 • Built with precision")
